from telegram import Update, ReplyKeyboardRemove
from telegram.ext import ContextTypes, ConversationHandler
from keyboards import *
from states import *
import traceback
import os

from pdf2docx import Converter

from PIL import Image
import pdfplumber
import fitz
from openpyxl import Workbook
import zipfile
from docx2pdf import convert
from docx import Document

import subprocess
from openpyxl import load_workbook
import csv


async def converter_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "📂 اختر  نوع الملف الذي تريد تحويله",
        reply_markup=convert_keyboard
    )
    return CONVERTER


async def converter_type(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()

    if query.data == "PDF":

        context.user_data["converter_type"] = "PDF"

        await query.edit_message_text(
            text="📄 اختر نوع التحويل الخاص بـ PDF:",
            reply_markup=convert_PDF
        )


    elif query.data == "Word":

        context.user_data["converter_type"] = "Word"

        await query.edit_message_text(
            text="📄 اختر نوع التحويل الخاص بـ Word:",
            reply_markup=convert_Word
        )


    elif query.data == "Excel":

        context.user_data["converter_type"] = "Excel"

        await query.edit_message_text(
            text="📄 اختر نوع التحويل الخاص بـ Excel:",
            reply_markup=convert_Excel
        )


    elif query.data == "Images":

        context.user_data["converter_type"] = "Images"

        await query.edit_message_text(
            text="📄 اختر نوع التحويل الخاص بـ Images:",
            reply_markup=convert_Images
        )
    elif query.data == "Compress":

        context.user_data["converter_type"] = "Compress"

        await query.edit_message_text(
            text="📄 اختر نوع التحويل الخاص بـ Compress:",
            reply_markup=convert_Compress
        )
    return CONVERTER_TYPE

async def send_document(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    context.user_data['conversion_type'] = query.data
    converter_type = context.user_data["converter_type"]

    if converter_type == "PDF":
        message = "📄 أرسل ملف PDF"

    elif converter_type == "Word":
        message = "📝 أرسل ملف Word"

    elif converter_type == "Excel":
        message = "📊 أرسل ملف Excel"

    elif converter_type == "Images":
        message = '''
            🖼 ارسل الصورة ك document
            اضغط على 📎 (زر إرفاق الملفات).
            لا تضغط على "Gallery" مباشرة لإرسال الصورة كصورة.
            اختر File (ملف) أو Document (مستند).
            انتقل إلى مجلد الصور في هاتفك.
            اختر الصورة (PNG أو JPG أو WEBP).
            أرسلها.
        '''

    elif converter_type == "Compress":
        message = "🗜 أرسل الملف الذي تريد ضغطه"

    await query.edit_message_text(message)
    return CONVERTE_FILE



async def convert_document(update: Update, context: ContextTypes.DEFAULT_TYPE):
    converter_type = context.user_data["converter_type"]
    conversion_type = context.user_data['conversion_type']


    await update.message.reply_text(
            "انتظر..."
    )
    

    document = update.message.document
    file_id = document.file_id
    telegram_file = await context.bot.get_file(file_id)
    filename = document.file_name
    await telegram_file.download_to_drive(f"uploads/{filename}")



    if converter_type == 'Images':
        try:
            if(conversion_type == 'PNG → JPG'):
                if(filename.lower().endswith(".png")):
                    image = Image.open(f"uploads/{filename}")
                    image = image.convert("RGB")
                    image.save(f"converted/photo{file_id}.jpg")
                    with open(f"converted/photo{file_id}.jpg", "rb") as f:
                        await update.message.reply_document(document=f)
                    await update.message.reply_text(
                        "اختر خدمة أخرى:",
                        reply_markup=reply_markup
                    )
                else:
                    await update.message.reply_text("❌ لقد اخترت تحويل PNG → JPG، لكن الملف المرسل ليس بصيغة PNG.")
            elif(conversion_type == 'JPG → PNG'):
                if(filename.lower().endswith((".jpg",".jpeg"))):
                    image = Image.open(f"uploads/{filename}")
                    image.save(f"converted/photo{file_id}.png")
                    with open(f"converted/photo{file_id}.png", "rb") as f:
                        await update.message.reply_document(document=f)
                    await update.message.reply_text(
                        "اختر خدمة أخرى:",
                        reply_markup=reply_markup
                    )
                else:
                    await update.message.reply_text("❌ لقد اخترت تحويل JPG → PNG، لكن الملف المرسل ليس بصيغة JPG.")
            elif(conversion_type == 'WEBP → JPG'):
                if(filename.lower().endswith(".webp")):
                    image = Image.open(f"uploads/{filename}")
                    image = image.convert("RGB")
                    image.save(f"converted/photo{file_id}.jpg")
                    with open(f"converted/photo{file_id}.jpg", "rb") as f:
                        await update.message.reply_document(document=f)
                    await update.message.reply_text(
                        "اختر خدمة أخرى:",
                        reply_markup=reply_markup
                    )
                else:
                    await update.message.reply_text("❌ لقد اخترت تحويل WEBP → JPG، لكن الملف المرسل ليس بصيغة WEBP.")

            elif(conversion_type == 'JPG → WEBP'):
                if(filename.lower().endswith((".jpg",".jpeg"))):
                    image = Image.open(f"uploads/{filename}")
                    image = image.convert("RGB")
                    image.save(f"converted/photo{file_id}.webp")
                    with open(f"converted/photo{file_id}.webp", "rb") as f:
                        await update.message.reply_document(document=f)
                    await update.message.reply_text(
                        "اختر خدمة أخرى:",
                        reply_markup=reply_markup
                    )
                else:
                    await update.message.reply_text("❌ لقد اخترت تحويل JPG → WEBP، لكن الملف المرسل ليس بصيغة JPG.")
            elif(conversion_type == 'Images → PDF'):
                if(filename.lower().endswith((".jpg",".jpeg",".png",".webp"))):
                    image = Image.open(f"uploads/{filename}")
                    image = image.convert("RGB")
                    image.save(f"converted/photo{file_id}.pdf")
                    with open(f"converted/photo{file_id}.pdf", "rb") as f:
                        await update.message.reply_document(document=f)
                    await update.message.reply_text(
                        "اختر خدمة أخرى:",
                        reply_markup=reply_markup
                    )
                else:
                    await update.message.reply_text("❌ لقد اخترت تحويل Images → PDF، لكن الملف المرسل ليس بصيغة PNG or JPG or WEBP.")
            
        except Exception:
            traceback.print_exc()
            await update.message.reply_text(
                "المرجو التأكد من الصورة المرسلة",
                reply_markup=reply_markup
            )
        finally:
            if os.path.exists(f"uploads/{filename}"):
                os.remove(f"uploads/{filename}")

            if os.path.exists(f"converted/photo{file_id}.jpg"):
                os.remove(f"converted/photo{file_id}.jpg")
    elif converter_type == 'PDF':
        if conversion_type == 'PDF → Word':
            if filename.lower().endswith(".pdf"):
                pdf_file = f"uploads/{filename}"
                word_file = f"converted/{file_id}.docx"
                cv = Converter(pdf_file)
                cv.convert(word_file)
                cv.close()
                with open(word_file, "rb") as f:
                    await update.message.reply_document(document=f)
                await update.message.reply_text(
                        "اختر خدمة أخرى:",
                        reply_markup=reply_markup
                    )
            else:
                await update.message.reply_text(
                    "❌ الملف المرسل ليس PDF."
                )
        elif conversion_type == 'PDF → Text':
            if filename.lower().endswith(".pdf"):
  
                text_file = f"converted/{file_id}.txt"

                with pdfplumber.open(f"uploads/{filename}") as pdf:

                    with open(text_file, "w", encoding="utf-8") as file:

                        for page in pdf.pages:

                            text = page.extract_text()

                            if text:

                                file.write(text)
                                file.write("\n\n")

                with open(text_file, "rb") as file:

                    await update.message.reply_document(document=file)

            else:

                await update.message.reply_text(
                    "❌ الملف ليس PDF."
                )
        
        elif conversion_type == 'PDF → Images':
            if filename.lower().endswith(".pdf"):
  
                pdf = fitz.open(f"uploads/{filename}")

                for page_number in range(len(pdf)):

                    page = pdf.load_page(page_number)

                    pix = page.get_pixmap()

                    image_path = f"converted/page_{page_number + 1}.png"

                    pix.save(image_path)

                    with open(image_path, "rb") as image:
                        await update.message.reply_document(document=image)

                pdf.close()
            else:

                await update.message.reply_text(
                    "❌ الملف ليس PDF."
                )
        elif conversion_type == 'PDF → Excel':
            if filename.lower().endswith(".pdf"):
                workbook = Workbook()

                sheet = workbook.active

                with pdfplumber.open(f"uploads/{filename}") as pdf:

                    for page in pdf.pages:

                        table = page.extract_table()

                        if table:

                            for row in table:

                                sheet.append(row)

                excel_file = f"converted/{file_id}.xlsx"

                workbook.save(excel_file)

                with open(excel_file, "rb") as file:

                    await update.message.reply_document(document=file)
            else:
                await update.message.reply_text(
                    "❌ الملف ليس PDF."
                )
    elif converter_type == 'Compress':
        if conversion_type == "Compress PDF":
            if filename.lower().endswith(".pdf"):
                with zipfile.ZipFile(f"converted/{file_id}.zip","w") as zip_file:
                    zip_file.write(f"uploads/{filename}")
                
                with open(f"converted/{file_id}.zip","rb") as file:
                    await update.message.reply_document(
                        document=file
                    )
            else:
                await update.message.reply_text(
                    "❌ الملف المرسل ليس PDF."
                )
        elif conversion_type == "Compress Images":
            if filename.lower().endswith((".png", ".jpg", ".jpeg", ".webp")):
                with zipfile.ZipFile(f"converted/{file_id}.zip","w") as zip_file:
                    zip_file.write(f"uploads/{filename}")
                with open(f"converted/{file_id}.zip","rb") as file:
                    await update.message.reply_document(
                        document=file
                    )
            else:
                await update.message.reply_text(
                    "❌ الملف المرسل ليس PDF."
                )
        elif conversion_type == "Compress Word":
            if filename.lower().endswith((".doc", ".docx")):

                with zipfile.ZipFile(f"converted/word{file_id}.zip","w") as zip_file:
                    zip_file.write(f"uploads/{filename}")

                with open(f"converted/word{file_id}.zip","rb") as file:
                    await update.message.reply_document(
                        document=file,
                    )


            else:

                await update.message.reply_text(
                    "❌ الملف المرسل ليس Word."
                )

        elif conversion_type == "Compress Excel":
            if filename.lower().endswith((".xls", ".xlsx")):

                with zipfile.ZipFile(f"converted/excel{file_id}.zip","w") as zip_file:
                    zip_file.write(f"uploads/{filename}")

                with open(f"converted/excel{file_id}.zip","rb") as file:

                    await update.message.reply_document(
                        document=file,
                    )


            else:

                await update.message.reply_text(
                    "❌ الملف المرسل ليس Excel."
                )

    elif converter_type == "Word":
        if conversion_type == "Word → PDF":
            if filename.lower().endswith((".doc",".docx")):
                convert(f"uploads/{filename}","converted/")
                with open(f"converted/{filename.replace('.docx','.pdf')}","rb") as file:
                    await update.message.reply_document(
                        document=file,
                        caption="📝 تم تحويل Word إلى PDF"
                    )
            else:
                await update.message.reply_text(
                    "❌ الملف المرسل ليس Word."
                )
        elif conversion_type == "Word → Text":
            if filename.lower().endswith(".docx"):
                doc = Document(f"uploads/{filename}")
                text = ""
                for paragraph in doc.paragraphs:
                    text += paragraph.text + "\n"
                txt_name = filename.replace(".docx",".txt")
                with open(f"converted/{txt_name}","w",encoding="utf-8") as file:
                    file.write(text)

                with open(f"converted/{txt_name}","rb") as file:
                    await update.message.reply_document(
                        document=file,
                        caption="📝 تم تحويل Word إلى Text"
                    )


            else:
                await update.message.reply_text(
                    "❌ المرجو إرسال ملف Word فقط"
                )

    elif converter_type == "Excel":
        if conversion_type == 'Excel → PDF':
            if filename.lower().endswith((".xls", ".xlsx")):
                subprocess.run([
                    r"C:\Program Files\LibreOffice\program\soffice.exe",
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    "converted",
                    f"uploads/{filename}"
                ])
                pdf_name = filename.replace(".xlsx", ".pdf")
                with open(f"converted/{pdf_name}", "rb") as file:
                    await update.message.reply_document(
                        document=file,
                        caption="📊 تم تحويل Excel إلى PDF"
                    )
                
            else:
                await update.message.reply_text(
                    "❌ المرجو إرسال ملف Excel فقط"
                )

        if conversion_type == 'Excel → CSV':
            if filename.lower().endswith((".xls", ".xlsx")):
                workbook = load_workbook(f"uploads/{filename}")
                sheet = workbook.active
                csv_name = os.path.splitext(filename)[0] + ".csv"   
                with open(f"converted/{csv_name}","w",newline="",encoding="utf-8") as file:
                    writer = csv.writer(file)
                    for row in sheet.iter_rows(values_only=True):
                        writer.writerow(row)

                with open(f"converted/{csv_name}", "rb") as file:
                    await update.message.reply_document(
                        document=file,
                        caption="📊 تم تحويل Excel إلى CSV"
                    )
                
            else:
                await update.message.reply_text(
                    "❌ المرجو إرسال ملف Excel فقط"
                )

# C:\Users\soufi\Videos











    
    return ConversationHandler.END


async def photo_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "❌ الرجاء إرسال الصورة كـ Document وليس كصورة عادية."
    )
    return CONVERTE_FILE